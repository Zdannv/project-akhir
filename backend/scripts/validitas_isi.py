"""Validitas isi kuesioner dengan penilaian ahli (indeks Aiken's V).

Butir RIASEC disusun penulis, butir BFI-10 diterjemahkan penulis, dan butir
minat bidang TI ditulis penulis. Sebelum dipakai untuk mengambil data, tiap
butir dinilai oleh beberapa ahli (dosen psikologi / bimbingan konseling /
praktisi rekrutmen TI) pada dua aspek, skala 1–5:

  * Kesesuaian  seberapa tepat butir mengukur dimensi yang dituju
  * Kejelasan   seberapa mudah kalimatnya dipahami responden

  python scripts/validitas_isi.py buat
      -> docs/validasi/lembar_validitas_isi.xlsx   (satu salinan per ahli)
  python scripts/validitas_isi.py hitung ahli1.xlsx ahli2.xlsx [ahli3.xlsx ...]
      -> docs/hasil-validitas-isi.md

Aiken's V = jumlah(nilai - 1) / (n x 4). p-value dihitung eksak terhadap
hipotesis bahwa ahli memilih nilai secara acak, jadi tidak bergantung pada
tabel kritis. Butir dengan p >= 0,05 pada aspek kesesuaian ditandai untuk
diperbaiki.
"""

from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from app import kuesioner as q  # noqa: E402
from app import statistik as st  # noqa: E402
from app.kb import AKAR, muat_kb  # noqa: E402

FOLDER = AKAR / "docs" / "validasi"
LEMBAR = FOLDER / "lembar_validitas_isi.xlsx"
KELUARAN = AKAR / "docs" / "hasil-validitas-isi.md"
DESKRIPSI_RIASEC = {
    "R": "Realistic: menyukai kegiatan praktis dengan alat, mesin, atau benda",
    "I": "Investigative: menyukai kegiatan menganalisis, meneliti, dan memecahkan masalah",
    "A": "Artistic: menyukai kegiatan kreatif dan ekspresif yang tidak terstruktur",
    "S": "Social: menyukai kegiatan membantu, mengajar, atau bekerja bersama orang lain",
    "E": "Enterprising: menyukai kegiatan memimpin, membujuk, dan berwirausaha",
    "C": "Conventional: menyukai kegiatan yang teratur, mengikuti prosedur, dan mengolah data",
}
DESKRIPSI_BIG5 = {
    "O": "Openness: keterbukaan terhadap pengalaman, imajinasi, dan seni",
    "C": "Conscientiousness: keteraturan, ketekunan, dan tanggung jawab",
    "E": "Extraversion: kecenderungan bersosialisasi dan energik",
    "A": "Agreeableness: keramahan, kepercayaan, dan kerja sama",
    "N": "Neuroticism: kecenderungan cemas dan mudah tertekan",
}


def butir() -> list[tuple[str, str, str, str]]:
    """(kode, instrumen, dimensi yang dituju, kalimat)."""
    hasil = [(k, "Minat RIASEC", DESKRIPSI_RIASEC[t], t_) for k, t, t_ in q.BUTIR_RIASEC]
    hasil += [(k, "Kepribadian (BFI-10)", DESKRIPSI_BIG5[d] + (" — butir berarah terbalik" if tb else ""),
               f"{q.AWALAN_BIG5.rstrip(' …')} {t_}") for k, d, tb, t_ in q.BUTIR_BIG5]
    hasil += [(f"BD_{b['kode']}", "Minat bidang TI", f"Minat pada bidang kerja: {b['nama']}", b["pertanyaan"])
              for b in muat_kb().bidang]
    return hasil


def buat(_: argparse.Namespace) -> None:
    FOLDER.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Petunjuk"
    for i, t in enumerate([
        "Lembar Validitas Isi Kuesioner", "",
        "Lembar 'Penilaian' berisi butir kuesioner yang dipakai sistem untuk mengukur minat dan kepribadian.",
        "Untuk tiap butir, beri dua nilai 1–5:",
        "  Kolom E  Kesesuaian: seberapa tepat butir ini mengukur dimensi pada kolom C",
        "           (1 = sangat tidak sesuai, 5 = sangat sesuai)",
        "  Kolom F  Kejelasan: seberapa mudah kalimatnya dipahami pelajar atau mahasiswa",
        "           (1 = sangat tidak jelas, 5 = sangat jelas)",
        "  Kolom G  Saran perbaikan kalimat (opsional)", "",
        "Butir minat dijawab responden dengan skala 'sangat tidak suka' sampai 'sangat suka';",
        "butir kepribadian dengan 'sangat tidak setuju' sampai 'sangat setuju'.", "",
        "Nama ahli:", "Bidang keahlian:", "Instansi:",
    ], 1):
        ws0.cell(i, 1, t).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    ws0.column_dimensions["A"].width = 100
    ws = wb.create_sheet("Penilaian")
    kepala = ["Kode", "Instrumen", "Dimensi yang diukur", "Butir", "Kesesuaian (1-5)", "Kejelasan (1-5)", "Saran"]
    ws.append(kepala)
    for c in range(1, len(kepala) + 1):
        ws.cell(1, c).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, c).fill = PatternFill("solid", fgColor="5B3FD6")
    daftar = butir()
    for b in daftar:
        ws.append([*b, None, None, None])
    for kol, lebar in zip("ABCDEFG", (9, 20, 46, 60, 16, 15, 40)):
        ws.column_dimensions[kol].width = lebar
    for baris in ws.iter_rows(min_row=2):
        for sel in baris:
            sel.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:F{len(daftar) + 1}")
    ws.freeze_panes = "A2"
    wb.save(LEMBAR)
    print(f"{len(daftar)} butir -> {LEMBAR.relative_to(AKAR)}")


def hitung(a: argparse.Namespace) -> None:
    if len(a.berkas) < 2:
        sys.exit("butuh lembar dari minimal dua ahli")
    nilai: dict[str, dict[str, list[int]]] = {}
    teks = {}
    for b in a.berkas:
        for kode, instr, _, kalimat, sesuai, jelas, _ in load_workbook(b, data_only=True)["Penilaian"].iter_rows(
                min_row=2, values_only=True):
            if not kode:
                continue
            teks[kode] = (instr, kalimat)
            for aspek, x in (("kesesuaian", sesuai), ("kejelasan", jelas)):
                if x is not None:
                    nilai.setdefault(kode, {}).setdefault(aspek, []).append(int(x))
    n_ahli = len(a.berkas)
    L = ["# Hasil Validitas Isi Kuesioner (Aiken's V)", "",
         f"Dibuat otomatis oleh `backend/scripts/validitas_isi.py` pada {date.today():%d-%m-%Y}.", "",
         f"{n_ahli} ahli, skala 1–5. V = jumlah(nilai − 1) / (n × 4). p-value eksak terhadap penilaian acak; "
         "butir dengan p ≥ 0,05 pada kesesuaian ditandai **perlu diperbaiki**.", "",
         "| Kode | Instrumen | Butir | V kesesuaian | p | V kejelasan | Keputusan |", "|---|---|---|---:|---:|---:|---|"]
    lemah = 0
    for kode, (instr, kalimat) in teks.items():
        v = nilai.get(kode, {})
        if len(v.get("kesesuaian", [])) < n_ahli:
            L.append(f"| {kode} | {instr} | {kalimat} | – | – | – | belum dinilai semua ahli |")
            continue
        s = st.aiken_v(v["kesesuaian"])
        j = st.aiken_v(v["kejelasan"]) if len(v.get("kejelasan", [])) == n_ahli else None
        keputusan = "valid" if s["p"] < 0.05 else "**perlu diperbaiki**"
        lemah += s["p"] >= 0.05
        vj = "–" if j is None else f"{j['v']:.2f}"
        L.append(f"| {kode} | {instr} | {kalimat} | {s['v']:.2f} | {s['p']:.3f} | {vj} | {keputusan} |")
    L += ["", f"{lemah} butir perlu diperbaiki. Dengan ahli sedikit, V harus tinggi untuk signifikan "
          "(mis. dua ahli baru signifikan bila keduanya memberi nilai 5), jadi gunakan minimal tiga ahli.", ""]
    KELUARAN.write_text("\n".join(L), encoding="utf-8")
    print("\n".join(L))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="aksi", required=True)
    sub.add_parser("buat")
    h = sub.add_parser("hitung")
    h.add_argument("berkas", nargs="+")
    a = ap.parse_args()
    buat(a) if a.aksi == "buat" else hitung(a)


if __name__ == "__main__":
    main()
