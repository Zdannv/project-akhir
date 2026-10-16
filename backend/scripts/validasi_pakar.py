"""Validasi pakar untuk relasi profesi–skill di basis pengetahuan.

66% relasi profesi–skill adalah kurasi penulis. Skrip ini menyiapkan lembar
penilaian untuk 2–3 pakar (praktisi TI atau dosen) dan menghitung tingkat
kesepakatan serta proporsi relasi yang dinilai relevan.

  python scripts/validasi_pakar.py buat [--per-profesi 6] [--seed 7]
      -> docs/validasi/lembar_validasi_skill.xlsx   (dikirim ke tiap pakar, satu salinan per orang)
         docs/validasi/kunci_validasi_skill.csv     (JANGAN dikirim; menyimpan asal tiap relasi)

  python scripts/validasi_pakar.py hitung pakar1.xlsx pakar2.xlsx [pakar3.xlsx ...]
      -> docs/hasil-validasi-pakar.md

Sampel diambil berstrata per profesi, separuh dari relasi hasil lowongan dan
separuh dari relasi kurasi bila tersedia, lalu diacak. Asal relasi tidak
ditampilkan ke pakar supaya penilaiannya tidak terpengaruh. Relasi dari data
lowongan berfungsi sebagai pembanding: kalau relasi kurasi dinilai relevan
dengan proporsi yang setara, kurasinya dapat dipertanggungjawabkan.
"""

from __future__ import annotations

import argparse
import csv
import random
import sys
from datetime import date
from itertools import combinations
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from app import statistik as st  # noqa: E402
from app.kb import AKAR, muat_kb  # noqa: E402

FOLDER = AKAR / "docs" / "validasi"
LEMBAR = FOLDER / "lembar_validasi_skill.xlsx"
KUNCI = FOLDER / "kunci_validasi_skill.csv"
KELUARAN = AKAR / "docs" / "hasil-validasi-pakar.md"
YA_TIDAK = ("Ya", "Tidak")

PETUNJUK = [
    "Lembar Validasi Relasi Profesi dan Skill",
    "",
    "Terima kasih telah bersedia menjadi validator. Lembar ini dipakai untuk menilai apakah skill yang",
    "dikaitkan sistem dengan sebuah profesi memang dibutuhkan oleh profesi tersebut di Indonesia.",
    "",
    "Cara mengisi (lembar 'Penilaian'):",
    "1. Kolom E 'Relevan?': pilih Ya kalau skill tersebut memang dibutuhkan atau lazim diminta untuk profesi",
    "   itu; pilih Tidak kalau tidak dibutuhkan, keliru, atau terlalu jauh dari pekerjaan profesi itu.",
    "2. Kolom F 'Seharusnya wajib?': pilih Ya kalau tanpa skill itu seseorang sulit diterima bekerja di",
    "   profesi tersebut; pilih Tidak kalau skill itu hanya nilai tambah. Isi kolom ini hanya bila kolom E = Ya.",
    "3. Kolom G 'Catatan': opsional, misalnya nama skill yang lebih tepat.",
    "",
    "Mohon tidak berdiskusi dengan validator lain sebelum selesai mengisi, supaya penilaiannya independen.",
    "Isi nama dan latar belakang Anda di kolom B pada tiga baris di bawah ini.",
    "",
    "Nama validator:",
    "Pekerjaan / jabatan:",
    "Lama pengalaman di bidang TI (tahun):",
]


def buat(a: argparse.Namespace) -> None:
    kb = muat_kb()
    nama_skill = kb.nama_skill()
    acak = random.Random(a.seed)
    sampel = []
    for p in kb.profesi:
        rel = kb.skill_profesi(p["id"])
        lowongan = [r for r in rel if r["sumber"] == "lowongan"]
        kurasi = [r for r in rel if r["sumber"] == "kurasi"]
        separuh = a.per_profesi // 2
        ambil_l = acak.sample(lowongan, min(separuh, len(lowongan)))
        ambil_k = acak.sample(kurasi, min(a.per_profesi - len(ambil_l), len(kurasi)))
        for r in ambil_l + ambil_k:
            sampel.append((p, r))
    acak.shuffle(sampel)

    FOLDER.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Petunjuk"
    for i, baris in enumerate(PETUNJUK, 1):
        ws0.cell(i, 1, baris).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    ws0.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Penilaian")
    kepala = ["No", "Profesi", "Deskripsi singkat profesi", "Skill", "Relevan?", "Seharusnya wajib?", "Catatan"]
    ws.append(kepala)
    for c in range(1, len(kepala) + 1):
        ws.cell(1, c).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, c).fill = PatternFill("solid", fgColor="5B3FD6")
    with open(KUNCI, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["no", "profesi_id", "skill_id", "sumber", "wajib_di_sistem"])
        for i, (p, r) in enumerate(sampel, 1):
            ws.append([i, p["nama"], p["deskripsi"], nama_skill[r["skill_id"]], None, None, None])
            w.writerow([i, p["id"], r["skill_id"], r["sumber"], int(r["wajib"])])
    for kol, lebar in zip("ABCDEFG", (6, 28, 60, 30, 11, 17, 40)):
        ws.column_dimensions[kol].width = lebar
    for baris in ws.iter_rows(min_row=2):
        for sel in baris:
            sel.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"Ya,Tidak"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:F{len(sampel) + 1}")
    ws.freeze_panes = "A2"
    wb.save(LEMBAR)
    n_k = sum(1 for _, r in sampel if r["sumber"] == "kurasi")
    print(f"{len(sampel)} relasi ({n_k} kurasi, {len(sampel) - n_k} lowongan) -> {LEMBAR.relative_to(AKAR)}")
    print(f"Kunci (jangan dikirim ke pakar) -> {KUNCI.relative_to(AKAR)}")


def baca(berkas: str) -> tuple[str, dict[int, tuple[str | None, str | None]]]:
    wb = load_workbook(berkas, data_only=True)
    nama = Path(berkas).stem
    for r in wb["Petunjuk"].iter_rows():
        if r[0].value and str(r[0].value).startswith("Nama validator"):
            ditulis = (r[1].value if len(r) > 1 else None) or str(r[0].value).split(":", 1)[1]
            nama = str(ditulis).strip() or nama
    isi = {}
    for no, _, _, _, rel, wajib, _ in wb["Penilaian"].iter_rows(min_row=2, values_only=True):
        if no is None:
            continue
        norm = lambda x: str(x).strip().capitalize() if x else None  # noqa: E731
        isi[int(no)] = (norm(rel), norm(wajib))
    return nama, isi


def hitung(a: argparse.Namespace) -> None:
    if len(a.berkas) < 2:
        sys.exit("butuh lembar dari minimal dua pakar")
    kunci = {int(r["no"]): r for r in csv.DictReader(open(KUNCI, encoding="utf-8"))}
    pakar = [baca(b) for b in a.berkas]
    pakar = [(f"{n} ({i})" if [m for m, _ in pakar].count(n) > 1 else n, isi) for i, (n, isi) in enumerate(pakar, 1)]
    nomor = sorted(n for n in kunci if all(isi.get(n, (None,))[0] in YA_TIDAK for _, isi in pakar))
    if not nomor:
        sys.exit("tidak ada butir yang dinilai lengkap oleh semua pakar")
    nilai = {nama: [isi[n][0] for n in nomor] for nama, isi in pakar}
    nama = list(nilai)

    L = ["# Hasil Validasi Pakar: Relasi Profesi–Skill", "",
         f"Dibuat otomatis oleh `backend/scripts/validasi_pakar.py` pada {date.today():%d-%m-%Y}.", "",
         f"{len(pakar)} pakar menilai {len(nomor)} relasi secara lengkap (dari {len(kunci)} relasi sampel).", "",
         "## Kesepakatan antarpakar (kolom Relevan?)", "", "| Pasangan | Kesepakatan | Cohen's κ |", "|---|---:|---:|"]
    for x, y in combinations(nama, 2):
        sepakat = sum(1 for u, v in zip(nilai[x], nilai[y]) if u == v) / len(nomor)
        k = st.cohen_kappa(nilai[x], nilai[y])
        L.append(f"| {x} – {y} | {sepakat:.1%} | {'–' if k is None else f'{k:.2f}'} |")
    if len(nama) >= 3:
        fk = st.fleiss_kappa([[nilai[p][i] for p in nama] for i in range(len(nomor))])
        L += ["", f"Fleiss' κ seluruh pakar: **{'–' if fk is None else f'{fk:.2f}'}**"]
    L += ["", "κ di bawah 0 berarti kesepakatan lebih rendah daripada kebetulan; makin mendekati 1 makin sepakat.", ""]

    mayoritas = {n: sum(1 for p in nama if nilai[p][i] == "Ya") > len(nama) / 2 for i, n in enumerate(nomor)}
    L += ["## Proporsi relasi yang dinilai relevan (suara mayoritas)", "",
          "| Asal relasi | n | Relevan | 95% CI (Wilson) |", "|---|---:|---:|---:|"]
    for sumber in ("kurasi", "lowongan"):
        ns = [n for n in nomor if kunci[n]["sumber"] == sumber]
        if ns:
            ya = sum(mayoritas[n] for n in ns)
            lo, hi = st.wilson(ya, len(ns))
            L.append(f"| {sumber} | {len(ns)} | {ya / len(ns):.1%} | {lo:.1%}–{hi:.1%} |")
    L += ["", "Relasi dari data lowongan menjadi pembanding. Kalau proporsi relasi kurasi setara, kurasi penulis "
          "dapat dipertanggungjawabkan; kalau jauh lebih rendah, relasi kurasi perlu ditinjau ulang.", ""]

    kb = muat_kb()
    nama_skill, nama_profesi = kb.nama_skill(), {p["id"]: p["nama"] for p in kb.profesi}
    ditolak = [n for n in nomor if not mayoritas[n]]
    L += ["## Relasi yang ditolak mayoritas pakar", ""]
    if ditolak:
        L += ["| No | Profesi | Skill | Asal |", "|---:|---|---|---|",
              *[f"| {n} | {nama_profesi[kunci[n]['profesi_id']]} | {nama_skill[kunci[n]['skill_id']]} | {kunci[n]['sumber']} |"
                for n in ditolak], "",
              "Relasi ini dikandidatkan untuk dihapus dari `data/sumber/pelengkap_profesi_skill.csv` "
              "(kurasi) atau diberi catatan (lowongan), lalu basis pengetahuan dibangun ulang.", ""]
    else:
        L += ["Tidak ada.", ""]
    KELUARAN.write_text("\n".join(L), encoding="utf-8")
    print("\n".join(L))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="aksi", required=True)
    b = sub.add_parser("buat")
    b.add_argument("--per-profesi", type=int, default=6)
    b.add_argument("--seed", type=int, default=7)
    h = sub.add_parser("hitung")
    h.add_argument("berkas", nargs="+")
    a = ap.parse_args()
    buat(a) if a.aksi == "buat" else hitung(a)


if __name__ == "__main__":
    main()
