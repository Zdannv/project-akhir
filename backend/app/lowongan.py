"""Lowongan nyata sebagai kasus uji di luar simulasi.

Persona simulasi dibangkitkan dari basis pengetahuan, sehingga mesin diuji
dengan data yang ia sendiri hasilkan. Di sini kasus ujinya adalah lowongan
LinkedIn mentah (April–Juli 2026): keahlian yang disebut di deskripsi lowongan
menjadi profil keahlian, dan profesi yang dimaksud judul lowongan menjadi
jawaban benar.

Keterbatasan yang dilaporkan bersama hasilnya: sebagian relasi keahlian
profesi data di basis pengetahuan diringkas dari lowongan yang sama, jadi uji
ini belum sepenuhnya lepas. Yang diuji adalah apakah profil keahlian tiap
profesi cukup membedakan satu lowongan dari lowongan profesi lain, pada teks
lowongan apa adanya.
"""

from __future__ import annotations

import re
from pathlib import Path

from .kb import BasisPengetahuan

# Urutan penting: pola yang lebih khusus diperiksa lebih dulu, dan judul yang
# tidak cocok dengan pola mana pun tidak dipakai (misalnya "data architect").
PETA_JUDUL = (
    ("C013", r"computer vision"),
    ("C014", r"\bnlp\b|natural language"),
    ("C008", r"business intelligence|\bbi\b|power ?bi|tableau developer"),
    ("C011", r"machine learning|\bml\b|\bmlops\b"),
    ("C012", r"\bai\b|artificial intelligence|\bllm\b|generative|agentic"),
    ("C010", r"data scien"),
    ("C009", r"data engineer|big data|\betl\b|data pipeline|analytics engineer"),
    ("C007", r"data analy|\banalytics\b|\banalyst\b"),
)


def profesi_dari_judul(judul: str) -> str | None:
    t = judul.lower()
    return next((pid for pid, pola in PETA_JUDUL if re.search(pola, t)), None)


# Alias yang di teks lowongan lebih sering bermakna lain: "make" sebagai kata
# kerja, "Apache" pada Apache Spark/Airflow, "query" dan "troubleshooting"
# sebagai kata umum, dan "C" pada "C-level" atau "Series C".
ALIAS_UMUM = {"make", "apache", "query", "c", "troubleshooting"}


def pola_skill(kb: BasisPengetahuan) -> list[tuple[str, re.Pattern]]:
    """Satu pola per keahlian dari nama dan aliasnya. Nama satu atau dua huruf
    (R, Go) dicocokkan peka huruf besar supaya tidak tertukar dengan kata biasa."""
    hasil = []
    for s in kb.skill:
        bagian = []
        for nama in [s["nama"], *s.get("alias", [])]:
            nama = nama.strip()
            if not nama or nama.lower() in ALIAS_UMUM:
                continue
            inti = re.escape(nama)
            bagian.append(f"(?-i:{inti})" if len(nama) <= 2 else inti)
        if bagian:
            pola = r"(?<![\w&/+#.])(?:" + "|".join(bagian) + r")(?![\w&+#])"
            hasil.append((s["id"], re.compile(pola, re.I)))
    return hasil


def skill_dari_teks(teks: str, pola: list[tuple[str, re.Pattern]]) -> set[str]:
    return {sid for sid, p in pola if p.search(teks)}


def baca_lowongan(berkas: Path, kb: BasisPengetahuan, minimal_skill: int = 3) -> list[dict]:
    """Lowongan unik (per tautan) yang judulnya terpetakan ke profesi dan
    deskripsinya menyebut paling sedikit `minimal_skill` keahlian."""
    import openpyxl

    ws = openpyxl.load_workbook(berkas, read_only=True).worksheets[0]
    kepala = [str(c) for c in next(ws.iter_rows(max_row=1, values_only=True))]
    pola = pola_skill(kb)
    terlihat, hasil = set(), []
    for baris in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(kepala, baris))
        tautan = str(r.get("link") or "").split("?")[0]
        if not tautan or tautan in terlihat or r.get("status_ekstraksi") != "completed":
            continue
        terlihat.add(tautan)
        pid = profesi_dari_judul(str(r.get("job_title") or ""))
        if not pid:
            continue
        skill = skill_dari_teks(f"{r.get('job_title') or ''}\n{r.get('job_description') or ''}", pola)
        if len(skill) < minimal_skill:
            continue
        hasil.append({"profesi_id": pid, "judul": str(r["job_title"]).strip(), "tautan": tautan,
                      "diunggah": str(r.get("posted_date") or "")[:10], "skill": sorted(skill)})
    return hasil
